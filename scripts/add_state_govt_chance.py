"""Add Raw Data column Z (Q11 'State Govt Chance') analysis to every analysis sheet,
mirroring the existing blocks/columns built for the neighbouring questions."""
import re, sys
from copy import copy
import openpyxl
from openpyxl.formula.tokenizer import Tokenizer, Token
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.formatting.rule import CellIsRule

SRC, DST = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC)

Z = "'Raw Data'!$Z$2:$Z$39997"
TITLE = "Q11. Should the State Government be given another chance, or do you want a change"
OPTS = [
    "Want to give another chance / ಇನ್ನೊಂದು ಅವಕಾಶ ಕೊಡಬೇಕು",
    "Want a change / ಬದಲಾವಣೆ ಬೇಕು",
    "Can't say / ಹೇಳಲು ಆಗಲ್ಲ",
]
REF = re.compile(r"^(\$?)([A-Z]{1,3})(\$?)(\d+)$")


def remap_formula(f, rowmap):
    """Rewrite row numbers of same-sheet refs in formula f via rowmap(row)->row."""
    if not (isinstance(f, str) and f.startswith("=")):
        return f
    tok = Tokenizer(f)
    out = []
    for t in tok.items:
        v = t.value
        if t.type == Token.OPERAND and t.subtype == Token.RANGE and "!" not in v:
            parts = []
            for p in v.split(":"):
                m = REF.match(p)
                if m:
                    p = f"{m[1]}{m[2]}{m[3]}{rowmap(int(m[4]))}"
                parts.append(p)
            v = ":".join(parts)
        out.append(v)
    return "=" + "".join(out)


def insert_rows(ws, at, n):
    """Insert n blank rows at `at`, keeping same-sheet refs, merges and row heights valid."""
    shift = lambda r: r + n if r >= at else r
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                c.value = remap_formula(c.value, shift)
    merges = [CellRange(str(m)) for m in ws.merged_cells.ranges]
    for m in merges:
        ws.unmerge_cells(str(m))
    heights = {r: d.height for r, d in ws.row_dimensions.items() if d.height is not None}
    ws.insert_rows(at, n)
    for m in merges:
        if m.min_row >= at:
            m.shift(row_shift=n)
        ws.merge_cells(m.coord)
    for r in list(ws.row_dimensions):
        if r >= at:
            ws.row_dimensions[r].height = None
    for r, h in heights.items():
        ws.row_dimensions[shift(r)].height = h


def copy_cell(src, dst, value):
    dst.value = value
    if src.has_style:
        dst._style = copy(src._style)


def clone_block(ws, s0, s_opt1, s_optN, s_total, d0):
    """Copy question block rows s0..s_total (Y block) to d0.., as a 3-option Z block.
    s_total may be None (block ends at last option row)."""
    head = s_opt1 - s0                       # rows before first option
    d_opt1 = d0 + head
    d_optN = d_opt1 + len(OPTS) - 1
    d_total = d_optN + 1 if s_total else None
    s_end = s_total or s_optN
    tmpl_text = ws.cell(s_opt1, 2).value

    def blockmap(r, cur_src, cur_dst):
        if r == cur_src:
            return cur_dst
        if s0 <= r < s_opt1:
            return d0 + (r - s0)
        if r == s_opt1:
            return d_opt1
        if r == s_optN:
            return d_optN
        if s_total and r == s_total:
            return d_total
        return r                              # outside block (e.g. caste header row 5)

    plan = [(s0 + i, d0 + i) for i in range(head)]
    plan += [(s_opt1, d_opt1 + k) for k in range(len(OPTS))]
    if s_total:
        plan.append((s_total, d_total))
    for sr, dr in plan:
        if ws.row_dimensions[sr].height is not None:
            ws.row_dimensions[dr].height = ws.row_dimensions[sr].height
        for c in range(1, ws.max_column + 1):
            src = ws.cell(sr, c)
            v = src.value
            if isinstance(v, str) and v.startswith("="):
                v = v.replace("'Raw Data'!$Y$2:$Y$39997", Z)
                v = remap_formula(v, lambda r: blockmap(r, sr, dr))
                if s_opt1 == sr:
                    k = dr - d_opt1
                    new = v.replace(f'"{tmpl_text}*"', f'"{OPTS[k]}*"')
                    assert new != v or "$Z$" not in v, (ws.title, src.coordinate, v[:120])
                    v = new
            elif sr == s0 and c == 1:
                v = TITLE
            elif sr == s_opt1 and c == 1:
                v = dr - d_opt1 + 1
            elif sr == s_opt1 and c == 2:
                v = OPTS[dr - d_opt1]
            copy_cell(src, ws.cell(dr, c), v)
    for m in list(ws.merged_cells.ranges):
        if s0 <= m.min_row < s_opt1:
            cr = CellRange(str(m))
            cr.shift(row_shift=d0 - s0)
            ws.merge_cells(cr.coord)
    return d_total or d_optN


# ---- Count: Y block rows 62..71 (title, hdr, base, 7 options) -> insert after row 72
ws = wb["Count"]
insert_rows(ws, 73, 7)
clone_block(ws, 62, 65, 71, None, 73)

# ---- Analysis %: Y block 72..83 (title, grp, hdr, base, 7 opts, total)
ws = wb["Analysis %"]
insert_rows(ws, 85, 9)
clone_block(ws, 72, 76, 82, 83, 85)

# ---- Analysis % Norm: Y block 82..95 (title, grp, hdr, base, base M, base F, 7 opts, total)
ws = wb["Analysis % Norm"]
insert_rows(ws, 97, 11)
clone_block(ws, 82, 88, 94, 95, 97)

# ---- Cross: new TABLE 6 = clone of TABLE 4 (rows 39..48, AV continue/stop) for Z
ws = wb["Cross"]
d0, s0 = 66, 39
av_opts = ["Should continue / ಮುಂದುವರಿಯಬೇಕು", "Should be stopped / ನಿಲ್ಲಿಸಬೇಕು", "Can't say / ಹೇಳಲು ಆಗಲ್ಲ"]
for i in range(10):
    sr, dr = s0 + i, d0 + i
    if ws.row_dimensions[sr].height is not None:
        ws.row_dimensions[dr].height = ws.row_dimensions[sr].height
    for c in range(1, ws.max_column + 1):
        src = ws.cell(sr, c)
        v = src.value
        if isinstance(v, str) and v.startswith("="):
            v = v.replace("'Raw Data'!$AV$2:$AV$39997", Z)
            for a, b in zip(av_opts, OPTS):
                v = v.replace(f'"{a}"', f'"{b}"')
            if c == 2:  # base: party respondents who answered Q11
                assert v.endswith(")")
                v = v[:-1] + f',{Z},"<>")'
            v = remap_formula(v, lambda r, sr=sr, dr=dr: r + (d0 - s0) if s0 <= r <= s0 + 9 else r)
        elif i == 0 and c == 1:
            v = ("TABLE 6: STATE GOVT — ANOTHER CHANCE vs WANT A CHANGE (Q11), broken down by PARTY the "
                 "respondent supports (rows = Party from Q3; columns = Q11 answer; % of row base). "
                 "Base = party supporters who answered Q11.")
        elif i == 1 and 3 <= c <= 5:
            v = OPTS[c - 3]
        copy_cell(src, ws.cell(dr, c), v)
for m in list(ws.merged_cells.ranges):
    if m.min_row == s0:
        cr = CellRange(str(m)); cr.shift(row_shift=d0 - s0); ws.merge_cells(cr.coord)


def add_ac_columns(ws, hdr_row, first, last, key_col, cols, tmpl_n, tmpl_pct, key_ref):
    """Append Z columns (base, another chance %, want change %, can't say %) to a per-row table."""
    c0 = cols
    names = ["State Govt Chance — Base (n)", "State Govt — Another Chance %",
             "State Govt — Want Change %", "State Govt Chance — Can't say %"]
    L = openpyxl.utils.get_column_letter
    for j, name in enumerate(names):
        col = c0 + j
        copy_cell(ws[f"{tmpl_n if j == 0 else tmpl_pct}{hdr_row}"], ws.cell(hdr_row, col), name)
        ws.column_dimensions[L(col)].width = max(ws.column_dimensions[tmpl_pct].width or 13, 16)
    bL = L(c0)
    for r in range(first, last + 1):
        crit = f"'Raw Data'!${key_col}$2:${key_col}$39997,{key_ref.format(r=r)}"
        copy_cell(ws[f"{tmpl_n}{r}"], ws.cell(r, c0), f'=COUNTIFS({crit},{Z},"<>")')
        for j, opt in enumerate(OPTS, start=1):
            copy_cell(ws[f"{tmpl_pct}{r}"], ws.cell(r, c0 + j),
                      f'=IFERROR(COUNTIFS({crit},{Z},"{opt}*")/${bL}{r},0)')
    return c0


# ---- Leader Perf.: AC-wise columns O..R
ws = wb["Leader Perf."]
c0 = add_ac_columns(ws, 1, 2, 69, "BN", 15, "D", "E", "$C{r}")
red = [cf for cf in ws.conditional_formatting if str(cf.sqref) == "F2:F69"][0].rules[0]
new = copy(red); new.priority = 99
ws.conditional_formatting.add("Q2:Q69", new)
# Overall grid: add a Q11 mini-grid under the existing performance grid
g = 81
copy_cell(ws["C71"], ws[f"C{g}"], "STATE GOVT — ANOTHER CHANCE vs WANT A CHANGE (Q11, % of respondents who answered Q11)")
ws.merge_cells(f"C{g}:E{g}")
copy_cell(ws["C73"], ws[f"C{g+2}"], "Answer"); copy_cell(ws["D73"], ws[f"D{g+2}"], "State Govt")
copy_cell(ws["C74"], ws[f"C{g+3}"], "Base (n)"); copy_cell(ws["D74"], ws[f"D{g+3}"], f'=COUNTIFS({Z},"<>")')
for k, (lab, opt) in enumerate(zip(["Another chance", "Want a change", "Can't say"], OPTS)):
    r = g + 4 + k
    copy_cell(ws[f"C{75+k}"], ws[f"C{r}"], lab)
    copy_cell(ws[f"D{75+k}"], ws[f"D{r}"], f'=IFERROR(COUNTIFS({Z},"{opt}*")/D{g+3},0)')

# ---- DW: date-wise columns R..U + TOTAL row 185
ws = wb["DW"]
c0 = add_ac_columns(ws, 3, 4, 183, "BO", 18, "B", "C", "$A{r}")
L = openpyxl.utils.get_column_letter
copy_cell(ws["B185"], ws.cell(185, c0), f'=COUNTIFS({Z},"<>")')
for j, opt in enumerate(OPTS, start=1):
    copy_cell(ws["C185"], ws.cell(185, c0 + j), f'=IFERROR(COUNTIFS({Z},"{opt}*")/${L(c0)}$185,0)')

wb.calculation.fullCalcOnLoad = True  # openpyxl keeps no cached values; make Excel recalc on open
wb.save(DST)
print("saved", DST)
